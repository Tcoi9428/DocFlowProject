from datetime import date, timedelta
from io import BytesIO, StringIO
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch
from zipfile import ZipFile
import xml.etree.ElementTree as ET

from django.contrib.auth.models import User
from django.core import mail
from django.core.management import call_command
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from django.utils import timezone
from openpyxl import Workbook as OpenpyxlWorkbook

from .correspondence import (
    attach_generated_memo_template,
    attach_generated_outgoing_template,
    build_outgoing_letter_template,
    build_service_memo_template,
    reserve_correspondence_number,
)
from .forms import DocumentForm
from .models import (
    ApprovalRoute,
    ApprovalStep,
    ApprovalTask,
    ContractKind,
    CorrespondenceDepartment,
    CorrespondenceRecord,
    CorrespondenceSequence,
    Document,
    DocumentApprover,
    DocumentComment,
    DocumentPurpose,
    EmailDelivery,
    Notification,
    DocumentType,
    PasswordResetRequest,
    RevisionRequest,
)
from .services import (
    approve_task,
    deliver_email,
    process_approval_reminders,
    resubmit_parallel_approval,
    return_for_revision,
    start_approval,
)


class DocumentWorkflowTests(TestCase):
    def setUp(self):
        self.author = User.objects.create_user(username="author", password="test")
        self.manager = User.objects.create_user(username="manager", password="test")
        self.director = User.objects.create_user(username="director", password="test")
        self.document_type = DocumentType.objects.create(name="Договоры", code="DOG")
        self.memo_type = DocumentType.objects.create(name="Служебные записки", code="SZ")
        self.order_type = DocumentType.objects.create(name="Приказы", code="PRK")
        ContractKind.objects.create(name="Договор поставки", code="SUPPLY")
        DocumentPurpose.objects.create(name="Основание для оплаты", code="PAYMENT")
        self.route = ApprovalRoute.objects.create(
            name="Типовой маршрут",
            document_type=self.document_type,
            is_default=True,
            route_type=ApprovalRoute.SEQUENTIAL,
        )
        ApprovalStep.objects.create(route=self.route, name="Руководитель", approver=self.manager, order=1)
        ApprovalStep.objects.create(route=self.route, name="Директор", approver=self.director, order=2)

    def test_system_number_is_generated_by_document_type_and_year(self):
        document = Document.objects.create(
            document_type=self.document_type,
            title="Поставка запчастей",
            author=self.author,
        )

        self.assertTrue(document.system_number.startswith("DOG-"))
        self.assertTrue(document.system_number.endswith("00001"))

    def test_sequential_approval_creates_next_task_and_approves_document(self):
        document = Document.objects.create(
            document_type=self.document_type,
            title="Ремонт оборудования",
            author=self.author,
            route=self.route,
        )

        with self.captureOnCommitCallbacks(execute=True):
            start_approval(document, self.author)
        first_task = ApprovalTask.objects.get(document=document, approver=self.manager)
        approve_task(first_task, self.manager, "OK")

        self.assertTrue(ApprovalTask.objects.filter(document=document, approver=self.director).exists())

        second_task = ApprovalTask.objects.get(document=document, approver=self.director)
        approve_task(second_task, self.director, "OK")
        document.refresh_from_db()

        self.assertEqual(document.status, Document.APPROVED)

    def test_create_form_keeps_selected_document_type_after_refresh(self):
        self.client.force_login(self.author)

        response = self.client.get(f"/documents/new/?document_type={self.document_type.id}")

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, f'value="{self.document_type.id}" selected')
        self.assertContains(response, f'name="document_type" value="{self.document_type.id}"')
        self.assertContains(response, "addEventListener")
        self.assertNotContains(response, "window.location")
        self.assertNotContains(response, "Обновить поля")

    def test_contract_fields_are_visible_only_for_contract_documents(self):
        contract_form = DocumentForm(initial={"document_type": self.document_type.id})
        memo_form = DocumentForm(initial={"document_type": self.memo_type.id})

        self.assertIn("contract_kind", contract_form.fields)
        self.assertIn("document_purpose", contract_form.fields)
        self.assertNotIn("contract_kind", memo_form.fields)
        self.assertNotIn("document_purpose", memo_form.fields)

    def test_order_form_hides_counterparty_and_amount(self):
        order_form = DocumentForm(initial={"document_type": self.order_type.id})

        self.assertNotIn("counterparty", order_form.fields)
        self.assertNotIn("amount", order_form.fields)

    def test_custom_document_approvers_drive_sequential_approval(self):
        document = Document.objects.create(
            document_type=self.document_type,
            title="Договор по пользовательскому маршруту",
            author=self.author,
            route=self.route,
        )
        DocumentApprover.objects.create(document=document, approver=self.manager, name="Первый этап", order=1)
        DocumentApprover.objects.create(document=document, approver=self.director, name="Второй этап", order=2)

        with self.captureOnCommitCallbacks(execute=True):
            start_approval(document, self.author)
        first_task = ApprovalTask.objects.get(document=document, approver=self.manager)

        approve_task(first_task, self.manager, "OK")

        self.assertTrue(ApprovalTask.objects.filter(document=document, approver=self.director).exists())

    def test_user_can_choose_parallel_mode_for_custom_approvers(self):
        document = Document.objects.create(
            document_type=self.document_type,
            title="Параллельный пользовательский маршрут",
            author=self.author,
            responsible=self.author,
            route=self.route,
            approval_route_type=ApprovalRoute.PARALLEL,
        )
        DocumentApprover.objects.create(document=document, approver=self.manager, order=1)
        DocumentApprover.objects.create(document=document, approver=self.director, order=2)

        start_approval(document, self.author)

        self.assertEqual(
            ApprovalTask.objects.filter(
                document=document,
                document_version=1,
                status=ApprovalTask.PENDING,
            ).count(),
            2,
        )

    def test_parallel_return_does_not_block_other_approvers(self):
        document = Document.objects.create(
            document_type=self.document_type,
            title="Независимое параллельное согласование",
            author=self.author,
            responsible=self.author,
            route=self.route,
            approval_route_type=ApprovalRoute.PARALLEL,
        )
        DocumentApprover.objects.create(document=document, approver=self.manager, order=1)
        DocumentApprover.objects.create(document=document, approver=self.director, order=2)
        start_approval(document, self.author)
        manager_task = ApprovalTask.objects.get(document=document, approver=self.manager)
        director_task = ApprovalTask.objects.get(document=document, approver=self.director)

        return_for_revision(manager_task, self.manager, comment="Исправить цену")
        document.refresh_from_db()
        director_task.refresh_from_db()

        self.assertEqual(document.status, Document.ON_APPROVAL)
        self.assertEqual(director_task.status, ApprovalTask.PENDING)
        self.assertTrue(
            RevisionRequest.objects.filter(
                document=document,
                requested_by=self.manager,
                status=RevisionRequest.OPEN,
            ).exists()
        )

        return_for_revision(director_task, self.director, comment="Уточнить срок")
        document.refresh_from_db()

        self.assertEqual(document.status, Document.RETURNED)
        self.assertEqual(document.revision_requests.filter(status=RevisionRequest.OPEN).count(), 2)

    def test_parallel_resubmit_targets_only_participants_with_comments(self):
        document = Document.objects.create(
            document_type=self.document_type,
            title="Выборочное повторное согласование",
            author=self.author,
            responsible=self.author,
            route=self.route,
            approval_route_type=ApprovalRoute.PARALLEL,
        )
        DocumentApprover.objects.create(document=document, approver=self.manager, order=1)
        DocumentApprover.objects.create(document=document, approver=self.director, order=2)
        start_approval(document, self.author)
        manager_task = ApprovalTask.objects.get(document=document, approver=self.manager)
        director_task = ApprovalTask.objects.get(document=document, approver=self.director)
        approve_task(manager_task, self.manager, "Согласовано")
        return_for_revision(director_task, self.director, comment="Добавить срок поставки")
        revision_request = document.revision_requests.get(status=RevisionRequest.OPEN)

        resubmit_parallel_approval(
            document,
            self.author,
            {revision_request.pk: "Срок поставки добавлен в раздел 4."},
        )
        document.refresh_from_db()
        revision_request.refresh_from_db()

        self.assertEqual(document.version, 2)
        self.assertEqual(document.status, Document.ON_APPROVAL)
        self.assertEqual(revision_request.status, RevisionRequest.RESOLVED)
        self.assertEqual(revision_request.resolved_in_version, 2)
        self.assertFalse(
            ApprovalTask.objects.filter(
                document=document,
                approver=self.manager,
                document_version=2,
                status=ApprovalTask.PENDING,
            ).exists()
        )
        retry_task = ApprovalTask.objects.get(
            document=document,
            approver=self.director,
            document_version=2,
            status=ApprovalTask.PENDING,
        )
        self.assertTrue(
            Notification.objects.filter(
                recipient=self.manager,
                document=document,
                title__icontains="версии 2",
            ).exists()
        )

        approve_task(retry_task, self.director, "Замечание устранено")
        document.refresh_from_db()

        self.assertEqual(document.status, Document.APPROVED)

    def test_parallel_revision_page_shows_separate_correction_fields(self):
        document = Document.objects.create(
            document_type=self.document_type,
            title="Два замечания к документу",
            author=self.author,
            responsible=self.author,
            route=self.route,
            approval_route_type=ApprovalRoute.PARALLEL,
        )
        DocumentApprover.objects.create(document=document, approver=self.manager, order=1)
        DocumentApprover.objects.create(document=document, approver=self.director, order=2)
        start_approval(document, self.author)
        return_for_revision(
            ApprovalTask.objects.get(document=document, approver=self.manager),
            self.manager,
            comment="Исправить цену",
        )
        return_for_revision(
            ApprovalTask.objects.get(document=document, approver=self.director),
            self.director,
            comment="Уточнить срок",
        )
        revision_requests = list(document.revision_requests.order_by("id"))
        self.client.force_login(self.author)

        response = self.client.get(f"/documents/{document.pk}/")

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Замечания к версии 1")
        self.assertContains(response, "Исправить цену")
        self.assertContains(response, "Уточнить срок")
        for revision_request in revision_requests:
            self.assertContains(response, f'name="correction_{revision_request.pk}"')

    def test_approval_recovers_missing_configured_approver_link(self):
        document = Document.objects.create(
            document_type=self.document_type,
            title="Legacy custom approval",
            author=self.author,
        )
        DocumentApprover.objects.create(document=document, approver=self.manager, order=1)
        DocumentApprover.objects.create(document=document, approver=self.director, order=2)
        start_approval(document, self.author)
        first_task = ApprovalTask.objects.get(document=document, approver=self.manager)
        first_task.configured_approver = None
        first_task.save(update_fields=["configured_approver", "updated_at"])

        approve_task(first_task, self.manager, "OK")

        self.assertTrue(
            ApprovalTask.objects.filter(
                document=document,
                approver=self.director,
                status=ApprovalTask.PENDING,
            ).exists()
        )

    def test_approval_finishes_orphan_task_without_route(self):
        document = Document.objects.create(
            document_type=self.document_type,
            title="Legacy approval without route",
            author=self.author,
            status=Document.ON_APPROVAL,
        )
        task = ApprovalTask.objects.create(document=document, approver=self.manager)

        approve_task(task, self.manager, "OK")
        document.refresh_from_db()

        self.assertEqual(document.status, Document.APPROVED)

    def test_start_approval_creates_notification_for_approver(self):
        document = Document.objects.create(
            document_type=self.document_type,
            title="Договор с уведомлением",
            author=self.author,
            route=self.route,
        )
        DocumentApprover.objects.create(document=document, approver=self.manager, name="Первый этап", order=1)

        with self.captureOnCommitCallbacks(execute=True):
            start_approval(document, self.author)

        self.assertTrue(
            Notification.objects.filter(
                recipient=self.manager,
                document=document,
                notification_type=Notification.APPROVAL_REQUIRED,
                is_read=False,
            ).exists()
        )

    @override_settings(
        DOCFLOW_EMAIL_SEND_IMMEDIATELY=False,
        DOCFLOW_BASE_URL="http://10.110.53.17:8010",
    )
    def test_approval_notification_queues_email_with_document_link(self):
        self.manager.email = "manager@example.com"
        self.manager.save(update_fields=["email"])
        document = Document.objects.create(
            document_type=self.document_type,
            title="Договор с email-уведомлением",
            author=self.author,
            route=self.route,
        )

        with self.captureOnCommitCallbacks(execute=True):
            start_approval(document, self.author)

        delivery = EmailDelivery.objects.get(notification__document=document)
        self.assertEqual(delivery.recipient_email, "manager@example.com")
        self.assertEqual(delivery.status, EmailDelivery.PENDING)
        self.assertEqual(delivery.link_url, f"http://10.110.53.17:8010/documents/{document.pk}/")

    @override_settings(
        EMAIL_BACKEND="django.core.mail.backends.locmem.EmailBackend",
        DOCFLOW_EMAIL_SEND_IMMEDIATELY=False,
        DOCFLOW_BASE_URL="http://10.110.53.17:8010",
    )
    def test_email_delivery_is_sent_and_marked_as_sent(self):
        self.manager.email = "manager@example.com"
        self.manager.save(update_fields=["email"])
        document = Document.objects.create(
            document_type=self.document_type,
            title="Проверка отправки email",
            author=self.author,
            route=self.route,
        )
        with self.captureOnCommitCallbacks(execute=True):
            start_approval(document, self.author)
        delivery = EmailDelivery.objects.get(notification__document=document)

        result = deliver_email(delivery.pk)

        delivery.refresh_from_db()
        self.assertTrue(result)
        self.assertEqual(delivery.status, EmailDelivery.SENT)
        self.assertEqual(delivery.attempts, 1)
        self.assertEqual(len(mail.outbox), 1)
        self.assertIn(f"/documents/{document.pk}/", mail.outbox[0].body)

    @override_settings(
        DOCFLOW_EMAIL_SEND_IMMEDIATELY=False,
        DOCFLOW_BASE_URL="http://10.110.53.17:8010",
    )
    def test_approval_reminder_is_created_once_for_current_due_date(self):
        self.manager.email = "manager@example.com"
        self.manager.save(update_fields=["email"])
        document = Document.objects.create(
            document_type=self.document_type,
            title="Документ с близким сроком",
            author=self.author,
            route=self.route,
        )
        start_approval(document, self.author)
        task = ApprovalTask.objects.get(document=document, approver=self.manager)
        task.due_date = timezone.localdate() + timedelta(days=1)
        task.save(update_fields=["due_date", "updated_at"])

        with self.captureOnCommitCallbacks(execute=True):
            first_count = process_approval_reminders()
        second_count = process_approval_reminders()

        task.refresh_from_db()
        reminder = Notification.objects.get(
            recipient=self.manager,
            document=document,
            notification_type=Notification.APPROVAL_REMINDER,
        )
        self.assertEqual(first_count, 1)
        self.assertEqual(second_count, 0)
        self.assertEqual(task.reminder_due_date, task.due_date)
        self.assertTrue(hasattr(reminder, "email_delivery"))
        self.assertIn(f"/documents/{document.pk}/", reminder.email_delivery.link_url)

    @override_settings(DOCFLOW_EMAIL_SEND_IMMEDIATELY=False)
    def test_responsible_is_notified_when_all_approvers_finish(self):
        responsible = User.objects.create_user(
            username="responsible",
            password="test",
            email="responsible@example.com",
        )
        document = Document.objects.create(
            document_type=self.document_type,
            title="Документ для ответственного",
            author=self.author,
            responsible=responsible,
            route=self.route,
        )
        start_approval(document, self.author)
        approve_task(ApprovalTask.objects.get(document=document, approver=self.manager), self.manager)
        with self.captureOnCommitCallbacks(execute=True):
            approve_task(ApprovalTask.objects.get(document=document, approver=self.director), self.director)

        notification = Notification.objects.get(
            recipient=responsible,
            document=document,
            title__icontains="согласован всеми участниками",
        )
        self.assertTrue(hasattr(notification, "email_delivery"))

    @override_settings(DOCFLOW_EMAIL_SEND_IMMEDIATELY=False)
    def test_responsible_receives_revision_email_with_actor_and_link(self):
        self.director.email = "director@example.com"
        self.director.save(update_fields=["email"])
        document = Document.objects.create(
            document_type=self.document_type,
            title="Документ на доработку",
            author=self.author,
            responsible=self.director,
            route=self.route,
        )
        start_approval(document, self.author)
        task = ApprovalTask.objects.get(document=document, approver=self.manager)

        with self.captureOnCommitCallbacks(execute=True):
            return_for_revision(task, self.manager, comment="Уточнить условия")

        notification = Notification.objects.get(
            recipient=self.director,
            document=document,
            title__icontains="отправлен на доработку",
        )
        self.assertIn("Уточнить условия", notification.message)
        self.assertTrue(hasattr(notification, "email_delivery"))
        self.assertIn(f"/documents/{document.pk}/", notification.email_delivery.link_url)

    @override_settings(DOCFLOW_EMAIL_SEND_IMMEDIATELY=False)
    def test_email_queue_failure_does_not_rollback_approval(self):
        self.author.email = "author@example.com"
        self.author.save(update_fields=["email"])
        document = Document.objects.create(
            document_type=self.document_type,
            title="Согласование независимо от почты",
            author=self.author,
            route=self.route,
        )
        start_approval(document, self.author)
        task = ApprovalTask.objects.get(document=document, approver=self.manager)

        with patch(
            "documents.services.queue_notification_email_by_id",
            side_effect=RuntimeError("Email queue is unavailable"),
        ):
            with self.captureOnCommitCallbacks(execute=True):
                approve_task(task, self.manager, comment="Согласовано")

        task.refresh_from_db()
        self.assertEqual(task.status, ApprovalTask.APPROVED)

    def test_return_for_revision_keeps_document_responsible_and_saves_comment(self):
        document = Document.objects.create(
            document_type=self.document_type,
            title="Договор на доработку",
            author=self.author,
            responsible=self.director,
            route=self.route,
        )

        start_approval(document, self.author)
        task = ApprovalTask.objects.get(document=document, approver=self.manager)
        return_for_revision(task, self.manager, comment="Исправить условия оплаты")
        document.refresh_from_db()

        self.assertEqual(document.status, Document.RETURNED)
        self.assertEqual(document.responsible, self.director)
        self.assertEqual(document.revision_requested_by, self.manager)
        self.assertEqual(document.revision_comment, "Исправить условия оплаты")
        self.assertTrue(
            DocumentComment.objects.filter(
                document=document,
                author=self.manager,
                text__icontains="Исправить условия оплаты",
            ).exists()
        )

    def test_resubmit_after_revision_increments_version_and_restarts_route(self):
        document = Document.objects.create(
            document_type=self.document_type,
            title="Повторное согласование",
            author=self.author,
            responsible=self.director,
            route=self.route,
        )

        start_approval(document, self.author)
        task = ApprovalTask.objects.get(document=document, approver=self.manager)
        return_for_revision(task, self.manager, comment="Нужна новая версия")

        self.client.force_login(self.director)
        response = self.client.post(
            f"/documents/{document.pk}/resubmit/",
            {"corrections": "Обновлен файл договора и условия оплаты."},
        )
        document.refresh_from_db()

        self.assertEqual(response.status_code, 302)
        self.assertEqual(document.version, 2)
        self.assertEqual(document.status, Document.ON_APPROVAL)
        self.assertTrue(
            ApprovalTask.objects.filter(
                document=document,
                approver=self.manager,
                status=ApprovalTask.PENDING,
            ).exists()
        )
        self.assertTrue(
            DocumentComment.objects.filter(
                document=document,
                author=self.director,
                text__icontains="Обновлен файл договора",
            ).exists()
        )

    def test_password_reset_request_notifies_admin_with_code(self):
        admin = User.objects.create_user(username="admin2", password="admin2", is_staff=True)
        response = self.client.post("/password-reset/", {"username": self.author.username})
        reset_request = PasswordResetRequest.objects.get(user=self.author)

        self.assertEqual(response.status_code, 302)
        self.assertEqual(len(reset_request.code), 4)
        self.assertTrue(reset_request.code.isdigit())
        self.assertTrue(
            Notification.objects.filter(
                recipient=admin,
                notification_type=Notification.PASSWORD_RESET,
                message__icontains=reset_request.code,
                link_url=f"/password-reset/admin/{reset_request.pk}/",
            ).exists()
        )

    def test_password_reset_confirm_changes_password_with_admin_code(self):
        reset_request = PasswordResetRequest.create_for_user(self.author)

        response = self.client.post(
            "/password-reset/confirm/",
            {
                "username": self.author.username,
                "code": reset_request.code,
                "new_password1": "pass1",
                "new_password2": "pass1",
            },
        )
        self.author.refresh_from_db()
        reset_request.refresh_from_db()

        self.assertEqual(response.status_code, 302)
        self.assertTrue(self.author.check_password("pass1"))
        self.assertEqual(reset_request.status, PasswordResetRequest.USED)

    def test_password_reset_confirm_rejects_password_without_digits(self):
        reset_request = PasswordResetRequest.create_for_user(self.author)

        response = self.client.post(
            "/password-reset/confirm/",
            {
                "username": self.author.username,
                "code": reset_request.code,
                "new_password1": "pass",
                "new_password2": "pass",
            },
        )
        self.author.refresh_from_db()

        self.assertEqual(response.status_code, 200)
        self.assertFalse(self.author.check_password("pass"))

    @override_settings(DOCFLOW_MAINTENANCE_MODE=True)
    def test_maintenance_mode_returns_503(self):
        response = self.client.get("/")

        self.assertEqual(response.status_code, 503)
        self.assertEqual(response.headers["Retry-After"], "300")
        self.assertContains(response, "Технические работы", status_code=503)

    @override_settings(DEBUG=False, DOCFLOW_MAINTENANCE_MODE=False)
    def test_unknown_page_uses_custom_404(self):
        response = self.client.get("/definitely-missing-page/")

        self.assertEqual(response.status_code, 404)
        self.assertContains(response, "Страница не найдена", status_code=404)


class OutgoingCorrespondenceTests(TestCase):
    def setUp(self):
        self.media_directory = TemporaryDirectory()
        self.media_override = override_settings(MEDIA_ROOT=Path(self.media_directory.name))
        self.media_override.enable()
        self.addCleanup(self.media_override.disable)
        self.addCleanup(self.media_directory.cleanup)

        self.user = User.objects.create_user(
            username="registrar",
            password="test",
            first_name="Иван",
            last_name="Иванов",
        )
        self.other_user = User.objects.create_user(username="second", password="test")
        self.department, _ = CorrespondenceDepartment.objects.update_or_create(
            code="02",
            defaults={"name": "Инжиниринг", "is_active": True},
        )
        CorrespondenceSequence.objects.update_or_create(
            kind=CorrespondenceRecord.OUTGOING,
            defaults={"next_number": 657},
        )

    def test_reserved_number_is_reused_for_same_user_and_unique_for_another_user(self):
        first = reserve_correspondence_number(self.user, CorrespondenceRecord.OUTGOING)
        repeated = reserve_correspondence_number(self.user, CorrespondenceRecord.OUTGOING)
        second = reserve_correspondence_number(self.other_user, CorrespondenceRecord.OUTGOING)

        self.assertEqual(first.pk, repeated.pk)
        self.assertEqual(first.sequence_number, 657)
        self.assertEqual(first.reserved_number, "01-__-657")
        self.assertEqual(second.sequence_number, 658)

    def test_registration_form_shows_current_date_in_html_calendar(self):
        self.client.force_login(self.user)

        response = self.client.get("/correspondence/outgoing/new/")

        self.assertContains(
            response,
            f'value="{timezone.localdate():%Y-%m-%d}"',
            html=False,
        )
        self.assertContains(response, 'data-searchable-select="true"', count=2)
        self.assertNotContains(response, "data-select-filter")

    def test_word_template_contains_generated_registration_number(self):
        subject = "О согласовании поставки запасных частей & оборудования"
        addressee = 'ООО "Заказчик & Партнеры"'
        addressee_person = "Иван Иванович"
        payload = build_outgoing_letter_template(
            "01-02-657",
            subject,
            addressee,
            addressee_person,
        )

        with ZipFile(BytesIO(payload)) as document:
            document_xml = document.read("word/document.xml")
            header_xml = document.read("word/header2.xml").decode("utf-8")

        word_namespace = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
        root = ET.fromstring(document_xml)
        first_row = next(root.iter(word_namespace + "tr"))
        cells = first_row.findall(word_namespace + "tc")
        marker_text = "".join((node.text or "") for node in cells[0].iter(word_namespace + "t"))
        number_text = "".join((node.text or "") for node in cells[1].iter(word_namespace + "t"))

        self.assertEqual(marker_text, "№")
        self.assertEqual(number_text, "01-02-657")
        document_text = "".join((node.text or "") for node in root.iter(word_namespace + "t"))
        self.assertIn(subject, document_text)
        self.assertNotIn("О О согласовании", document_text)
        self.assertIn(addressee, document_text)
        self.assertIn(f"Уважаемый {addressee_person}!", document_text)
        self.assertNotIn("Наименование компании ХХ «ХХХ»", document_text)
        self.assertIn("Сервис-Инжиниринг", header_xml)

    def test_template_download_attaches_draft_to_reservation(self):
        record = reserve_correspondence_number(self.user, CorrespondenceRecord.OUTGOING)
        self.client.force_login(self.user)

        response = self.client.post(
            f"/correspondence/outgoing/{record.pk}/template/",
            {
                "department": self.department.pk,
                "subject": "О согласовании поставки запасных частей",
                "addressee": "ООО Заказчик",
                "addressee_person": "Иван Иванович",
            },
        )
        record.refresh_from_db()

        self.assertEqual(response.status_code, 200)
        self.assertEqual(record.registration_number, "01-02-657")
        self.assertTrue(record.draft_file)
        self.assertEqual(record.draft_original_name, "Исходящее письмо 01-02-657.docx")
        self.assertEqual(record.subject, "О согласовании поставки запасных частей")

    def test_template_download_reuses_subject_from_existing_draft(self):
        record = reserve_correspondence_number(self.user, CorrespondenceRecord.OUTGOING)
        record.department = self.department
        record.subject = "О ранее сформированном письме"
        record.addressee = "ООО Заказчик"
        record.addressee_person = "Иван Иванович"
        attach_generated_outgoing_template(record)
        self.client.force_login(self.user)

        response = self.client.post(
            f"/correspondence/outgoing/{record.pk}/template/",
            {"department": self.department.pk},
        )
        record.refresh_from_db()

        self.assertEqual(response.status_code, 200)
        self.assertEqual(record.subject, "О ранее сформированном письме")
        self.assertTrue(record.draft_file)

    def test_outgoing_letter_can_be_registered_without_signed_file(self):
        record = reserve_correspondence_number(self.user, CorrespondenceRecord.OUTGOING)
        self.client.force_login(self.user)

        response = self.client.post(
            "/correspondence/outgoing/new/",
            {
                "reservation_id": record.pk,
                "department": self.department.pk,
                "reply_to": "",
                "related_document_number": "ВХ-2025-104",
                "addressee": "ООО Заказчик",
                "addressee_person": "Директору Петрову П.П.",
                "subject": "О согласовании поставки запасных частей",
                "registration_date": "2026-09-01",
                "executor": self.user.pk,
            },
        )
        record.refresh_from_db()

        self.assertEqual(response.status_code, 302)
        self.assertEqual(record.status, CorrespondenceRecord.REGISTERED)
        self.assertEqual(record.registration_number, "01-02-657")
        self.assertEqual(record.related_document_number, "ВХ-2025-104")
        self.assertFalse(record.has_signed_document)

        registry_response = self.client.get("/correspondence/outgoing/?scope=all")
        self.assertContains(registry_response, "01-02-657")
        self.assertContains(registry_response, "Не прикреплен")

    def test_signed_document_can_be_added_after_registration(self):
        record = CorrespondenceRecord.objects.create(
            kind=CorrespondenceRecord.OUTGOING,
            status=CorrespondenceRecord.REGISTERED,
            sequence_number=657,
            registration_number="01-02-657",
            department=self.department,
            addressee="ООО Заказчик",
            addressee_person="Директору",
            subject="Письмо",
            executor=self.user,
            created_by=self.user,
        )
        self.client.force_login(self.user)
        uploaded_file = SimpleUploadedFile("signed.pdf", b"signed document", content_type="application/pdf")

        response = self.client.post(
            f"/correspondence/outgoing/{record.pk}/signed-file/",
            {"signed_file": uploaded_file},
        )
        record.refresh_from_db()

        self.assertEqual(response.status_code, 302)
        self.assertTrue(record.signed_file)
        self.assertEqual(record.signed_original_name, "signed.pdf")


class IncomingCorrespondenceTests(TestCase):
    def setUp(self):
        self.media_directory = TemporaryDirectory()
        self.media_override = override_settings(MEDIA_ROOT=Path(self.media_directory.name))
        self.media_override.enable()
        self.addCleanup(self.media_override.disable)
        self.addCleanup(self.media_directory.cleanup)

        self.user = User.objects.create_user(
            username="incoming_registrar",
            password="test",
            first_name="Петр",
            last_name="Петров",
        )
        self.other_user = User.objects.create_user(username="incoming_second", password="test")
        self.department, _ = CorrespondenceDepartment.objects.update_or_create(
            code="01",
            defaults={"name": "Общий отдел", "is_active": True},
        )
        CorrespondenceSequence.objects.update_or_create(
            kind=CorrespondenceRecord.INCOMING,
            defaults={"next_number": 25},
        )
        self.outgoing = CorrespondenceRecord.objects.create(
            kind=CorrespondenceRecord.OUTGOING,
            status=CorrespondenceRecord.REGISTERED,
            sequence_number=657,
            registration_number="01-01-657",
            department=self.department,
            addressee="ООО Заказчик",
            addressee_person="Директору",
            subject="О направлении документов",
            executor=self.user,
            created_by=self.user,
        )

    def test_incoming_number_is_reserved_independently_for_each_user(self):
        first = reserve_correspondence_number(self.user, CorrespondenceRecord.INCOMING)
        repeated = reserve_correspondence_number(self.user, CorrespondenceRecord.INCOMING)
        second = reserve_correspondence_number(self.other_user, CorrespondenceRecord.INCOMING)

        self.assertEqual(first.pk, repeated.pk)
        self.assertEqual(first.sequence_number, 25)
        self.assertEqual(first.reserved_number, "02-__-25")
        self.assertEqual(second.sequence_number, 26)

    def test_incoming_registration_form_shows_current_date(self):
        self.client.force_login(self.user)

        response = self.client.get("/correspondence/incoming/new/")

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Регистрация входящего письма")
        self.assertContains(response, f'value="{timezone.localdate():%Y-%m-%d}"', html=False)
        self.assertContains(response, "02-__-25")
        self.assertContains(response, 'data-searchable-select="true"', count=1)
        self.assertNotContains(response, "data-select-filter")

    def test_incoming_letter_can_be_registered_with_outgoing_link(self):
        record = reserve_correspondence_number(self.user, CorrespondenceRecord.INCOMING)
        self.client.force_login(self.user)
        uploaded_file = SimpleUploadedFile("received.pdf", b"received letter", content_type="application/pdf")

        response = self.client.post(
            "/correspondence/incoming/new/",
            {
                "reservation_id": record.pk,
                "department": self.department.pk,
                "subject": "Ответ на запрос документации",
                "sender": "ООО Заказчик",
                "related_outgoing": self.outgoing.pk,
                "related_document_number": "СТАРЫЙ-НОМЕР",
                "registration_date": "2026-09-02",
                "resolution": "Передать в общий отдел",
                "incoming_file": uploaded_file,
            },
        )
        record.refresh_from_db()

        self.assertRedirects(response, f"/correspondence/incoming/{record.pk}/")
        self.assertEqual(record.status, CorrespondenceRecord.REGISTERED)
        self.assertEqual(record.registration_number, "02-01-25")
        self.assertEqual(record.related_outgoing, self.outgoing)
        self.assertEqual(record.related_document_number, "")
        self.assertEqual(record.created_by, self.user)
        self.assertEqual(record.executor, self.user)
        self.assertTrue(record.incoming_file)
        self.assertEqual(record.incoming_original_name, "received.pdf")

        outgoing_response = self.client.get(f"/correspondence/outgoing/{self.outgoing.pk}/")
        self.assertContains(outgoing_response, "02-01-25")
        self.assertContains(outgoing_response, "Ответ на запрос документации")

    def test_manual_related_number_is_preserved_for_old_registry(self):
        record = reserve_correspondence_number(self.user, CorrespondenceRecord.INCOMING)
        self.client.force_login(self.user)

        response = self.client.post(
            "/correspondence/incoming/new/",
            {
                "reservation_id": record.pk,
                "department": self.department.pk,
                "subject": "Письмо из старого реестра",
                "sender": "АО Поставщик",
                "related_outgoing": "",
                "related_document_number": "ИСХ-2025-104",
                "registration_date": "2026-09-02",
                "resolution": "",
            },
        )
        record.refresh_from_db()

        self.assertEqual(response.status_code, 302)
        self.assertIsNone(record.related_outgoing)
        self.assertEqual(record.related_document_number, "ИСХ-2025-104")

    def test_incoming_file_can_be_attached_after_registration(self):
        record = CorrespondenceRecord.objects.create(
            kind=CorrespondenceRecord.INCOMING,
            status=CorrespondenceRecord.REGISTERED,
            sequence_number=25,
            registration_number="02-01-25",
            department=self.department,
            sender="ООО Заказчик",
            subject="Входящее письмо",
            executor=self.user,
            created_by=self.user,
        )
        self.client.force_login(self.user)
        uploaded_file = SimpleUploadedFile("incoming.pdf", b"incoming letter", content_type="application/pdf")

        response = self.client.post(
            f"/correspondence/incoming/{record.pk}/file/",
            {"incoming_file": uploaded_file},
        )
        record.refresh_from_db()

        self.assertRedirects(response, f"/correspondence/incoming/{record.pk}/")
        self.assertTrue(record.incoming_file)
        self.assertEqual(record.incoming_original_name, "incoming.pdf")

        registry_response = self.client.get("/correspondence/incoming/?scope=all")
        self.assertContains(registry_response, "02-01-25")
        self.assertContains(registry_response, "Прикреплен")


class MemoCorrespondenceTests(TestCase):
    def setUp(self):
        self.media_directory = TemporaryDirectory()
        self.media_override = override_settings(MEDIA_ROOT=Path(self.media_directory.name))
        self.media_override.enable()
        self.addCleanup(self.media_override.disable)
        self.addCleanup(self.media_directory.cleanup)

        self.user = User.objects.create_user(
            username="memo_registrar",
            password="test",
            first_name="Иван",
            last_name="Иванов",
        )
        self.other_user = User.objects.create_user(username="memo_second", password="test")
        self.department, _ = CorrespondenceDepartment.objects.update_or_create(
            code="02",
            defaults={"name": "Инжиниринг", "is_active": True},
        )
        CorrespondenceSequence.objects.update_or_create(
            kind=CorrespondenceRecord.MEMO,
            defaults={"next_number": 669},
        )

    def test_counter_can_be_lowered_and_reservation_uses_next_free_number(self):
        CorrespondenceRecord.objects.create(
            kind=CorrespondenceRecord.MEMO,
            status=CorrespondenceRecord.REGISTERED,
            sequence_number=687,
            registration_number="03-02-687",
            department=self.department,
            subject="Некорректная старая запись",
            executor=self.user,
            created_by=self.user,
        )
        sequence = CorrespondenceSequence.objects.get(kind=CorrespondenceRecord.MEMO)
        sequence.next_number = 669
        sequence.full_clean()
        sequence.save()

        first = reserve_correspondence_number(self.user, CorrespondenceRecord.MEMO)
        second = reserve_correspondence_number(self.other_user, CorrespondenceRecord.MEMO)

        self.assertEqual(first.sequence_number, 669)
        self.assertEqual(second.sequence_number, 670)

    def test_registration_form_shows_number_date_and_template_notes(self):
        self.client.force_login(self.user)

        response = self.client.get("/correspondence/memos/new/")

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Регистрация служебной записки")
        self.assertContains(response, "03-__-669")
        self.assertContains(response, f'value="{timezone.localdate():%Y-%m-%d}"', html=False)
        self.assertContains(response, "автоматически попадет", count=4)
        self.assertContains(response, 'data-searchable-select="true"', count=1)

    def test_word_template_contains_memo_fields(self):
        payload = build_service_memo_template(
            "03-02-669",
            date(2026, 9, 4),
            "Об организации рабочих мест & графика",
            "Иванов Иван Иванович",
        )

        with ZipFile(BytesIO(payload)) as document:
            document_xml = document.read("word/document.xml")

        word_namespace = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
        root = ET.fromstring(document_xml)
        rows = list(root.iter(word_namespace + "tr"))
        number_row_cells = rows[3].findall(word_namespace + "tc")
        number_text = "".join(
            (node.text or "") for node in number_row_cells[1].iter(word_namespace + "t")
        )
        date_text = "".join(
            (node.text or "") for node in number_row_cells[3].iter(word_namespace + "t")
        )
        document_text = "".join((node.text or "") for node in root.iter(word_namespace + "t"))

        self.assertEqual(number_text, "03-02-669")
        self.assertEqual(date_text, "04.09.2026")
        self.assertIn("Об организации рабочих мест & графика", document_text)
        self.assertIn("Иванов Иван Иванович", document_text)
        self.assertIn("Уважаемый Иванов Иван Иванович!", document_text)
        self.assertNotIn("О…", document_text)
        self.assertNotIn("…. ", document_text)

    def test_template_download_and_registration_keep_draft_current(self):
        record = reserve_correspondence_number(self.user, CorrespondenceRecord.MEMO)
        self.client.force_login(self.user)

        download_response = self.client.post(
            f"/correspondence/memos/{record.pk}/template/",
            {
                "department": self.department.pk,
                "registration_date": "2026-09-04",
                "subject": "О первоначальном вопросе",
                "addressee_person": "Петров Петр Петрович",
            },
        )
        record.refresh_from_db()

        self.assertEqual(download_response.status_code, 200)
        self.assertEqual(record.registration_number, "03-02-669")
        self.assertTrue(record.draft_file)
        self.assertEqual(record.draft_original_name, "Служебная записка 03-02-669.docx")

        registration_response = self.client.post(
            "/correspondence/memos/new/",
            {
                "reservation_id": record.pk,
                "department": self.department.pk,
                "registration_date": "2026-09-04",
                "subject": "Об актуальном вопросе",
                "addressee_person": "Сидоров Сидор Сидорович",
                "executor": self.user.pk,
                "resolution": "Проверено",
            },
        )
        record.refresh_from_db()

        self.assertRedirects(registration_response, f"/correspondence/memos/{record.pk}/")
        self.assertEqual(record.status, CorrespondenceRecord.REGISTERED)
        self.assertEqual(record.subject, "Об актуальном вопросе")
        self.assertFalse(record.has_signed_document)
        with record.draft_file.open("rb") as draft_file:
            with ZipFile(draft_file) as document:
                document_xml = document.read("word/document.xml").decode("utf-8")
        self.assertIn("Об актуальном вопросе", document_xml)
        self.assertIn("Сидоров Сидор Сидорович", document_xml)
        self.assertNotIn("первоначальном", document_xml)

        registry_response = self.client.get("/correspondence/memos/?scope=all")
        self.assertContains(registry_response, "03-02-669")
        self.assertContains(registry_response, "Не прикреплен")
        self.assertContains(registry_response, f"/correspondence/memos/{record.pk}/")

    def test_signed_document_can_be_added_after_memo_registration(self):
        record = CorrespondenceRecord.objects.create(
            kind=CorrespondenceRecord.MEMO,
            status=CorrespondenceRecord.REGISTERED,
            sequence_number=669,
            registration_number="03-02-669",
            department=self.department,
            subject="Служебная записка",
            addressee_person="Петров П.П.",
            executor=self.user,
            created_by=self.user,
        )
        self.client.force_login(self.user)
        uploaded_file = SimpleUploadedFile("memo-signed.pdf", b"signed", content_type="application/pdf")

        response = self.client.post(
            f"/correspondence/memos/{record.pk}/signed-file/",
            {"signed_file": uploaded_file},
        )
        record.refresh_from_db()

        self.assertRedirects(response, f"/correspondence/memos/{record.pk}/")
        self.assertTrue(record.signed_file)
        self.assertEqual(record.signed_original_name, "memo-signed.pdf")


class CorrespondenceImportTests(TestCase):
    def setUp(self):
        self.registrar = User.objects.create_user(
            username="archive_registrar",
            password="test",
            first_name="Иван",
            last_name="Петрикин",
        )
        self.executor = User.objects.create_user(
            username="executor",
            password="test",
            first_name="Ольга",
            last_name="Ауст",
        )
        profile = self.executor.userprofile
        profile.patronymic = "Владимировна"
        profile.save(update_fields=["patronymic", "updated_at"])
        for code, name in [("01", "Общий отдел"), ("02", "Инжиниринг"), ("03", "Сервис")]:
            CorrespondenceDepartment.objects.update_or_create(
                code=code,
                defaults={"name": name, "is_active": True},
            )
        for kind in (CorrespondenceRecord.OUTGOING, CorrespondenceRecord.INCOMING, CorrespondenceRecord.MEMO):
            CorrespondenceSequence.objects.update_or_create(kind=kind, defaults={"next_number": 1})

        self.temp_directory = TemporaryDirectory()
        self.addCleanup(self.temp_directory.cleanup)
        self.xlsx_path = Path(self.temp_directory.name) / "2026.xlsx"
        self.build_workbook(self.xlsx_path)

    @staticmethod
    def build_workbook(path):
        workbook = OpenpyxlWorkbook()
        outgoing = workbook.active
        outgoing.title = " 01-0х-хх регистрация исходящих"
        outgoing.append(["*Подсказка"])
        outgoing.append(
            ["Номер ИСХ", "В ответ на номер ВХД", "Адресат", "Кому", "Наименование", "Дата", "Столбец1", "Исполнитель"]
        )
        outgoing.append(["01-02-661", "", "АО Заказчик", "Директору", "Письмо 2026", date(2026, 9, 2), "", "Ауст О.В."])
        outgoing.append(["01-02-100", "", "АО Архив", "Директору", "Письмо 2025", date(2025, 9, 2), "", "Ауст О.В."])

        incoming = workbook.create_sheet("02-0х-хх регистрация входящих")
        incoming.append(
            ["№ п/п", "Номер исходящего в ответ на входящее ", "Дата регистрации", "Отправитель", "Номер документа", "Дата", "Наименование", "Резолюция", "Ответ"]
        )
        incoming.append([142, "02-01-240", date(2026, 8, 10), "АО Поставщик", "22/468-К", date(2026, 8, 6), "Входящее письмо", "Передать в работу", "01-02-661 от 02.09.2026"])

        memos = workbook.create_sheet("03-0х-хх регистрация СЗ")
        memos.append(["номер", "дата", "наименование", "на кого", "исполнитель", "Примечание"])
        memos.append(["03-03-514", date(2026, 9, 1), "О закупке ТМЦ", "Петрикину И.И.", "Неизвестный И.И.", "Согласовано"])
        workbook.save(path)

    def run_import(self, apply=False):
        output = StringIO()
        options = {
            "year": 2026,
            "registrar": self.registrar.username,
            "stdout": output,
        }
        if apply:
            options["apply"] = True
        call_command("import_correspondence_xlsx", str(self.xlsx_path), **options)
        return output.getvalue()

    def test_import_is_previewable_idempotent_and_preserves_historical_fields(self):
        preview = self.run_import()

        self.assertIn("ПРЕДВАРИТЕЛЬНАЯ ПРОВЕРКА", preview)
        self.assertEqual(CorrespondenceRecord.objects.count(), 0)

        result = self.run_import(apply=True)

        self.assertIn("ИМПОРТ ВЫПОЛНЕН", result)
        self.assertEqual(CorrespondenceRecord.objects.count(), 3)
        outgoing = CorrespondenceRecord.objects.get(kind=CorrespondenceRecord.OUTGOING)
        incoming = CorrespondenceRecord.objects.get(kind=CorrespondenceRecord.INCOMING)
        memo = CorrespondenceRecord.objects.get(kind=CorrespondenceRecord.MEMO)
        self.assertEqual(outgoing.registration_number, "01-02-661")
        self.assertEqual(outgoing.executor, self.executor)
        self.assertEqual(outgoing.legacy_executor_name, "Ауст О.В.")
        self.assertEqual(incoming.external_document_number, "22/468-К")
        self.assertEqual(incoming.document_date, date(2026, 8, 6))
        self.assertEqual(incoming.related_outgoing, outgoing)
        self.assertEqual(memo.executor, self.registrar)
        self.assertEqual(memo.legacy_executor_name, "Неизвестный И.И.")
        self.assertTrue(all(record.is_historical_import for record in (outgoing, incoming, memo)))
        self.assertGreater(outgoing.sequence_number, 1_000_000_000)
        self.assertEqual(CorrespondenceSequence.objects.get(kind=CorrespondenceRecord.OUTGOING).next_number, 662)
        self.assertEqual(CorrespondenceSequence.objects.get(kind=CorrespondenceRecord.INCOMING).next_number, 241)
        self.assertEqual(CorrespondenceSequence.objects.get(kind=CorrespondenceRecord.MEMO).next_number, 515)

        self.run_import(apply=True)

        self.assertEqual(CorrespondenceRecord.objects.count(), 3)
        next_record = reserve_correspondence_number(self.registrar, CorrespondenceRecord.OUTGOING)
        self.assertEqual(next_record.sequence_number, 662)

    def test_imported_memos_are_visible_in_registry(self):
        self.run_import(apply=True)
        self.client.force_login(self.registrar)

        response = self.client.get("/correspondence/memos/?scope=all")

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "03-03-514")
        self.assertContains(response, "О закупке ТМЦ")

    def test_registry_searches_correspondence_by_number_subject_and_party(self):
        self.run_import(apply=True)
        self.client.force_login(self.registrar)

        cases = (
            ("outgoing", "01-02-661", "01-02-661"),
            ("outgoing", "АО Заказчик", "Письмо 2026"),
            ("incoming", "22/468-К", "02-01-240"),
            ("incoming", "АО Поставщик", "Входящее письмо"),
            ("memos", "закупке ТМЦ", "03-03-514"),
            ("memos", "Петрикину", "О закупке ТМЦ"),
        )
        for section, query, expected in cases:
            with self.subTest(section=section, query=query):
                response = self.client.get(
                    f"/correspondence/{section}/",
                    {"scope": "all", "q": query},
                )
                self.assertEqual(response.status_code, 200)
                self.assertEqual(response.context["query"], query)
                self.assertEqual(len(response.context["records"]), 1)
                self.assertContains(response, expected)

        empty_response = self.client.get(
            "/correspondence/outgoing/",
            {"scope": "all", "q": "несуществующий документ"},
        )
        self.assertEqual(len(empty_response.context["records"]), 0)
        self.assertContains(empty_response, "ничего не найдено")

# Create your tests here.
