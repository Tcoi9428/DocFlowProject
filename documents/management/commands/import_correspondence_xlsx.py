import re
from collections import Counter, defaultdict
from datetime import date, datetime, time
from pathlib import Path

from django.contrib.auth import get_user_model
from django.core.management.base import BaseCommand, CommandError
from django.core.exceptions import ObjectDoesNotExist
from django.db import transaction
from django.utils import timezone
from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel

from documents.models import (
    CorrespondenceDepartment,
    CorrespondenceRecord,
    CorrespondenceSequence,
)


KIND_LABELS = {
    CorrespondenceRecord.OUTGOING: "Исходящие",
    CorrespondenceRecord.INCOMING: "Входящие",
    CorrespondenceRecord.MEMO: "Служебные записки",
}
EXPECTED_KIND_CODES = {
    CorrespondenceRecord.OUTGOING: "01",
    CorrespondenceRecord.INCOMING: "02",
    CorrespondenceRecord.MEMO: "03",
}
SHEET_PREFIXES = {
    CorrespondenceRecord.OUTGOING: "01-",
    CorrespondenceRecord.INCOMING: "02-",
    CorrespondenceRecord.MEMO: "03-",
}
NUMBER_PATTERN = re.compile(r"(?<!\d)(01|02|03)-(01|02|03)-(\d+)(?!\d)")
HISTORICAL_SEQUENCE_BASE = 1_000_000_000


def normalize_text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return " ".join(str(value).replace("\xa0", " ").split()).strip()


def normalize_header(value):
    return normalize_text(value).lower().replace("ё", "е")


def trim(value, max_length):
    return normalize_text(value)[:max_length]


def parse_date(value, workbook_epoch):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)):
        parsed = from_excel(value, workbook_epoch)
        return parsed.date() if isinstance(parsed, datetime) else parsed
    text_value = normalize_text(value)
    for date_format in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(text_value, date_format).date()
        except ValueError:
            continue
    return None


def extract_number(value, expected_kind=None):
    text_value = normalize_text(value)
    matches = list(NUMBER_PATTERN.finditer(text_value))
    if expected_kind:
        expected_code = EXPECTED_KIND_CODES[expected_kind]
        matches = [match for match in matches if match.group(1) == expected_code]
    if not matches:
        return None
    return matches[0].group(0)


def person_parts(value):
    text_value = normalize_text(value).lower().replace("ё", "е")
    tokens = re.findall(r"[a-zа-я]+", text_value, flags=re.IGNORECASE)
    if not tokens:
        return "", ""
    return tokens[0], "".join(token[0] for token in tokens[1:] if token)


class ImportStats:
    def __init__(self):
        self.scanned = 0
        self.year_rows = 0
        self.created = 0
        self.existing = 0
        self.missing_date = 0
        self.unmatched_executors = Counter()
        self.number_warnings = 0
        self.links_created = 0
        self.max_sequence = 0


class Command(BaseCommand):
    help = "Импортирует исторический реестр корреспонденции из XLSX. Без --apply выполняет проверку."

    def add_arguments(self, parser):
        parser.add_argument("xlsx_path", help="Полный путь к файлу XLSX")
        parser.add_argument("--year", type=int, default=timezone.localdate().year, help="Импортируемый год")
        parser.add_argument("--registrar", required=True, help="Логин пользователя-регистратора")
        parser.add_argument(
            "--kinds",
            nargs="+",
            choices=list(KIND_LABELS),
            default=list(KIND_LABELS),
            help="Разделы для импорта",
        )
        parser.add_argument("--apply", action="store_true", help="Сохранить результат в базе данных")

    def handle(self, *args, **options):
        source_path = Path(options["xlsx_path"]).expanduser().resolve()
        if not source_path.is_file():
            raise CommandError(f"Файл не найден: {source_path}")
        if source_path.suffix.lower() not in {".xlsx", ".xlsm"}:
            raise CommandError("Поддерживаются только файлы .xlsx и .xlsm.")

        User = get_user_model()
        try:
            registrar = User.objects.select_related("userprofile").get(username=options["registrar"])
        except User.DoesNotExist as exc:
            raise CommandError(f"Пользователь '{options['registrar']}' не найден.") from exc

        departments = {
            department.code: department
            for department in CorrespondenceDepartment.objects.filter(code__in=["01", "02", "03"])
        }
        missing_codes = sorted({"01", "02", "03"} - set(departments))
        if missing_codes:
            raise CommandError(f"Не найдены подразделения корреспонденции с кодами: {', '.join(missing_codes)}")

        try:
            workbook = load_workbook(source_path, read_only=True, data_only=True)
        except Exception as exc:
            raise CommandError(f"Не удалось открыть Excel: {exc}") from exc

        users_by_exact_key, users_by_short_key = self.build_user_indexes(User)
        stats = {kind: ImportStats() for kind in options["kinds"]}
        source_name = source_path.name
        pending_links = []

        try:
            with transaction.atomic():
                for kind in options["kinds"]:
                    worksheet = self.find_worksheet(workbook, kind)
                    self.import_worksheet(
                        workbook,
                        worksheet,
                        kind,
                        options["year"],
                        source_name,
                        registrar,
                        departments,
                        users_by_exact_key,
                        users_by_short_key,
                        stats[kind],
                        pending_links,
                    )

                self.resolve_links(pending_links, stats)
                self.update_sequences(stats)
                if not options["apply"]:
                    transaction.set_rollback(True)
        finally:
            workbook.close()

        mode = "ИМПОРТ ВЫПОЛНЕН" if options["apply"] else "ПРЕДВАРИТЕЛЬНАЯ ПРОВЕРКА"
        self.stdout.write(self.style.SUCCESS(mode))
        self.stdout.write(f"Файл: {source_path}")
        self.stdout.write(f"Год: {options['year']}")
        for kind in options["kinds"]:
            item = stats[kind]
            created_label = "Создано" if options["apply"] else "Будет создано"
            self.stdout.write("")
            self.stdout.write(self.style.MIGRATE_HEADING(KIND_LABELS[kind]))
            self.stdout.write(f"  Строк просмотрено: {item.scanned}")
            self.stdout.write(f"  Строк за выбранный год: {item.year_rows}")
            self.stdout.write(f"  {created_label}: {item.created}")
            self.stdout.write(f"  Уже существовало: {item.existing}")
            self.stdout.write(f"  Строк без распознаваемой даты: {item.missing_date}")
            self.stdout.write(f"  Нестандартных номеров: {item.number_warnings}")
            self.stdout.write(f"  Связей с другими письмами: {item.links_created}")
            self.stdout.write(f"  Следующий номер после импорта: {item.max_sequence + 1 if item.max_sequence else 'без изменений'}")
            if item.unmatched_executors:
                self.stdout.write(
                    self.style.WARNING(
                        "  Не сопоставлены исполнители: "
                        + ", ".join(
                            f"{name or '[не указан]'} ({count})"
                            for name, count in item.unmatched_executors.most_common(12)
                        )
                    )
                )
        if not options["apply"]:
            self.stdout.write("")
            self.stdout.write(self.style.WARNING("Изменения не сохранены. Для импорта повторите команду с --apply."))

    @staticmethod
    def build_user_indexes(User):
        exact = defaultdict(list)
        short = defaultdict(list)
        for user in User.objects.filter(is_active=True).select_related("userprofile"):
            surname = normalize_text(user.last_name).lower().replace("ё", "е")
            if not surname:
                continue
            initials = normalize_text(user.first_name)[:1].lower()
            try:
                initials += normalize_text(user.userprofile.patronymic)[:1].lower()
            except ObjectDoesNotExist:
                pass
            exact[(surname, initials)].append(user)
            short[(surname, initials[:1])].append(user)
        return exact, short

    @staticmethod
    def find_worksheet(workbook, kind):
        prefix = SHEET_PREFIXES[kind]
        for worksheet in workbook.worksheets:
            if normalize_text(worksheet.title).startswith(prefix):
                return worksheet
        raise CommandError(f"В книге не найден лист, начинающийся с '{prefix}'.")

    def import_worksheet(
        self,
        workbook,
        worksheet,
        kind,
        target_year,
        source_name,
        registrar,
        departments,
        users_by_exact_key,
        users_by_short_key,
        stats,
        pending_links,
    ):
        header_row, columns = self.find_headers(worksheet, kind)
        empty_run = 0
        seen_data = False
        max_column = max(columns.values()) + 1
        for row_number, row in enumerate(
            worksheet.iter_rows(
                min_row=header_row + 1,
                max_row=5000,
                max_col=max_column,
                values_only=True,
            ),
            start=header_row + 1,
        ):
            if not any(normalize_text(value) for value in row):
                if seen_data:
                    empty_run += 1
                    if empty_run >= 100:
                        break
                continue
            seen_data = True
            empty_run = 0
            stats.scanned += 1

            registration_date = parse_date(row[columns["registration_date"]], workbook.epoch)
            if not registration_date:
                stats.missing_date += 1
                continue
            if registration_date.year != target_year:
                continue
            stats.year_rows += 1

            registration_number = trim(row[columns["registration_number"]], 40)
            number_match = NUMBER_PATTERN.search(registration_number)
            if number_match and number_match.group(1) == EXPECTED_KIND_CODES[kind]:
                department_code = number_match.group(2)
                stats.max_sequence = max(stats.max_sequence, int(number_match.group(3)))
            else:
                department_code = number_match.group(2) if number_match else "01"
                stats.number_warnings += 1
            if not registration_number:
                registration_number = f"АРХИВ-{EXPECTED_KIND_CODES[kind]}-{target_year}-{row_number}"[:40]

            import_source = f"{source_name}:{worksheet.title}"
            if CorrespondenceRecord.objects.filter(
                kind=kind,
                import_source=import_source,
                import_source_row=row_number,
            ).exists():
                stats.existing += 1
                continue

            row_data = self.row_data(kind, row, columns, workbook.epoch)
            existing = CorrespondenceRecord.objects.filter(
                kind=kind,
                registration_number=registration_number,
                registration_date=registration_date,
                subject=row_data["subject"],
            ).first()
            if existing:
                stats.existing += 1
                pending_links.append((existing, kind, row_data["related_document_number"]))
                continue

            legacy_executor = row_data.pop("legacy_executor_name")
            executor = self.match_executor(
                legacy_executor,
                users_by_exact_key,
                users_by_short_key,
            )
            if executor is None:
                executor = registrar
                if kind != CorrespondenceRecord.INCOMING:
                    stats.unmatched_executors[legacy_executor] += 1

            registered_at = timezone.make_aware(datetime.combine(registration_date, time(12, 0)))
            record = CorrespondenceRecord.objects.create(
                kind=kind,
                status=CorrespondenceRecord.REGISTERED,
                sequence_number=self.next_historical_sequence(kind, target_year, row_number),
                registration_number=registration_number,
                department=departments.get(department_code, departments["01"]),
                executor=executor,
                created_by=registrar,
                reserved_at=registered_at,
                registered_at=registered_at,
                legacy_executor_name=legacy_executor,
                is_historical_import=True,
                import_source=import_source,
                import_source_row=row_number,
                **row_data,
            )
            stats.created += 1
            pending_links.append((record, kind, record.related_document_number))

    @staticmethod
    def find_headers(worksheet, kind):
        required = {
            CorrespondenceRecord.OUTGOING: {
                "registration_number": "номер исх",
                "related_document_number": "в ответ на номер вхд",
                "addressee": "адресат",
                "addressee_person": "кому",
                "subject": "наименование",
                "registration_date": "дата",
                "executor": "исполнитель",
            },
            CorrespondenceRecord.INCOMING: {
                "registration_number": "номер исходящего в ответ на входящее",
                "registration_date": "дата регистрации",
                "sender": "отправитель",
                "external_document_number": "номер документа",
                "document_date": "дата",
                "subject": "наименование",
                "resolution": "резолюция",
                "related_document_number": "ответ",
            },
            CorrespondenceRecord.MEMO: {
                "registration_number": "номер",
                "registration_date": "дата",
                "subject": "наименование",
                "addressee_person": "на кого",
                "executor": "исполнитель",
                "resolution": "примечание",
            },
        }[kind]
        for row_number, row in enumerate(
            worksheet.iter_rows(min_row=1, max_row=10, values_only=True),
            start=1,
        ):
            header_map = {normalize_header(value): index for index, value in enumerate(row) if normalize_header(value)}
            if all(header in header_map for header in required.values()):
                return row_number, {field: header_map[header] for field, header in required.items()}
        raise CommandError(f"Не удалось распознать заголовки листа '{worksheet.title}'.")

    @staticmethod
    def row_data(kind, row, columns, workbook_epoch):
        related_number = trim(row[columns.get("related_document_number", 0)], 100) if "related_document_number" in columns else ""
        if kind == CorrespondenceRecord.OUTGOING:
            return {
                "related_document_number": related_number,
                "addressee": trim(row[columns["addressee"]], 250),
                "addressee_person": trim(row[columns["addressee_person"]], 250),
                "subject": trim(row[columns["subject"]], 300) or "Без наименования (архивная запись)",
                "registration_date": parse_date(row[columns["registration_date"]], workbook_epoch),
                "legacy_executor_name": trim(row[columns["executor"]], 200),
            }
        if kind == CorrespondenceRecord.INCOMING:
            return {
                "related_document_number": related_number,
                "sender": trim(row[columns["sender"]], 250),
                "external_document_number": trim(row[columns["external_document_number"]], 100),
                "document_date": parse_date(row[columns["document_date"]], workbook_epoch),
                "subject": trim(row[columns["subject"]], 300) or "Без наименования (архивная запись)",
                "resolution": normalize_text(row[columns["resolution"]]),
                "registration_date": parse_date(row[columns["registration_date"]], workbook_epoch),
                "legacy_executor_name": "",
            }
        return {
            "related_document_number": "",
            "addressee_person": trim(row[columns["addressee_person"]], 250),
            "subject": trim(row[columns["subject"]], 300) or "Без наименования (архивная запись)",
            "resolution": normalize_text(row[columns["resolution"]]),
            "registration_date": parse_date(row[columns["registration_date"]], workbook_epoch),
            "legacy_executor_name": trim(row[columns["executor"]], 200),
        }

    @staticmethod
    def match_executor(value, users_by_exact_key, users_by_short_key):
        surname, initials = person_parts(value)
        if not surname:
            return None
        exact_matches = users_by_exact_key.get((surname, initials[:2]), [])
        if len(exact_matches) == 1:
            return exact_matches[0]
        short_matches = users_by_short_key.get((surname, initials[:1]), [])
        return short_matches[0] if len(short_matches) == 1 else None

    @staticmethod
    def next_historical_sequence(kind, year, row_number):
        candidate = HISTORICAL_SEQUENCE_BASE + (year - 2000) * 1_000_000 + row_number
        while CorrespondenceRecord.objects.filter(kind=kind, sequence_number=candidate).exists():
            candidate += 10_000
        return candidate

    @staticmethod
    def resolve_links(pending_links, stats):
        for record, kind, raw_number in pending_links:
            if not raw_number:
                continue
            if kind == CorrespondenceRecord.OUTGOING:
                related_kind = CorrespondenceRecord.INCOMING
                field_name = "reply_to"
            elif kind == CorrespondenceRecord.INCOMING:
                related_kind = CorrespondenceRecord.OUTGOING
                field_name = "related_outgoing"
            else:
                continue
            number = extract_number(raw_number, related_kind)
            if not number:
                continue
            related = CorrespondenceRecord.objects.filter(
                kind=related_kind,
                registration_number=number,
                status=CorrespondenceRecord.REGISTERED,
            ).order_by("-registration_date", "-id").first()
            if related:
                setattr(record, field_name, related)
                record.save(update_fields=[field_name, "updated_at"])
                stats[kind].links_created += 1

    @staticmethod
    def update_sequences(stats):
        for kind, item in stats.items():
            if not item.max_sequence:
                continue
            sequence, _ = CorrespondenceSequence.objects.get_or_create(kind=kind, defaults={"next_number": 1})
            desired_next = item.max_sequence + 1
            if sequence.next_number < desired_next:
                sequence.next_number = desired_next
                sequence.save(update_fields=["next_number", "updated_at"])
